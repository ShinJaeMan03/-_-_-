import csv
import io
import json
import os
import sys
import re
import subprocess
import tkinter as tk
import webbrowser
from tkinter import filedialog, messagebox, ttk

try:
	from PIL import Image, ImageTk
except ImportError:
	Image = None
	ImageTk = None

try:
	import fitz
except ImportError:
	try:
		import pymupdf as fitz
	except ImportError:
		fitz = None

try:
	import pytesseract
except ImportError:
	pytesseract = None

try:
	from openpyxl import Workbook, load_workbook
except ImportError:
	Workbook = None
	load_workbook = None


APP_TITLE = "중간과제 - 액상화 예비평가"
MAX_EVALUATION_DEPTH = 20.0


def parse_spt(value):
	"""14/30, 50/5 같은 표기에서 원시 타격횟수를 추출합니다."""
	match = re.search(r"(\d+(?:\.\d+)?)\s*/", value.strip())
	return float(match.group(1)) if match else None


def is_refusal_spt(value):
	"""관입불가로 보는 1/50, 5/50, 50/5 등의 표기를 찾습니다."""
	match = re.search(r"(\d+(?:\.\d+)?)\s*/\s*(\d+(?:\.\d+)?)", str(value).strip())
	if not match:
		return False
	first = float(match.group(1))
	second = float(match.group(2))
	return (second >= 50 and first <= 5) or (first >= 50 and second < 30)


def number(value, default=None):
	try:
		return float(str(value).strip())
	except (TypeError, ValueError):
		return default


def infer_spt_interval(end_depth, count):
	"""Excel 기준에 따라 측점 수를 1.0m 또는 1.5m 간격으로 해석합니다."""
	if count <= 0 or end_depth <= 0:
		return None
	if count * 1.0 <= end_depth < count * 1.5:
		return 1.0
	if count * 1.5 <= end_depth:
		return 1.5
	options = (1.0, 1.5)
	return min(options, key=lambda interval: abs(count - round(end_depth / interval)))


class LiquefactionApp(tk.Tk):
	def __init__(self):
		super().__init__()
		self.title(APP_TITLE)
		self.geometry("1450x850")
		self.minsize(1050, 650)
		self.rows = []
		self.image_path = ""
		self.preview_image = None
		self.preview_source = None
		self.preview_scale = 1.0
		self.preview_offset = [0, 0]
		self.preview_drag_start = None
		self.pdf_document = None
		self.pdf_page_index = 0
		self.evaluation_mode = "일반 예비평가"
		self.history_file = os.path.join(os.path.dirname(os.path.abspath(__file__)), "liquefaction_history.json")
		self.history = self.load_history()
		self._build_style()
		self._build_ui()
		self.add_row()

	def _build_style(self):
		style = ttk.Style(self)
		try:
			style.theme_use("clam")
		except tk.TclError:
			pass
		style.configure("Title.TLabel", font=("맑은 고딕", 18, "bold"))
		style.configure("Section.TLabelframe.Label", font=("맑은 고딕", 11, "bold"))
		style.configure("Treeview", rowheight=28, font=("맑은 고딕", 9), borderwidth=1, relief="solid")
		style.configure("Treeview.Heading", font=("맑은 고딕", 9, "bold"), relief="raised", borderwidth=1)
		style.configure("Result.TLabel", font=("맑은 고딕", 11, "bold"))

	def _build_ui(self):
		top = ttk.Frame(self, padding=(12, 10, 12, 5))
		top.pack(fill="x")
		ttk.Label(top, text=APP_TITLE, style="Title.TLabel").pack(side="left")
		ttk.Label(top, text="입력자료를 기준으로 지층별 예비평가 대상 여부를 확인합니다.").pack(side="left", padx=18, pady=5)
		tk.Button(top, text="평가 결과 저장", command=self.export_excel).pack(side="right", padx=2)
		tk.Button(top, text="기록 비교 저장", command=self.export_history_comparison).pack(side="right", padx=2)

		self.notebook = ttk.Notebook(self)
		self.notebook.pack(fill="both", expand=True, padx=12, pady=5)
		self.input_tab = ttk.Frame(self.notebook)
		self.image_tab = ttk.Frame(self.notebook)
		self.result_tab = ttk.Frame(self.notebook)
		self.history_tab = ttk.Frame(self.notebook)
		self.notebook.add(self.input_tab, text="시추자료 입력")
		self.notebook.add(self.image_tab, text="파일 자동 읽기")
		self.notebook.add(self.result_tab, text="예비평가 결과")
		self.notebook.add(self.history_tab, text="이전 기록/비교")
		self._build_input_tab()
		self._build_image_tab()
		self._build_result_tab()
		self._build_history_tab()

	def _build_input_tab(self):
		meta = ttk.LabelFrame(self.input_tab, text="기본 정보", style="Section.TLabelframe", padding=8)
		meta.pack(fill="x", padx=5, pady=5)
		
		# 1줄: 시추번호, 조사명, 위치, 표고
		fields_row1 = [("시추번호", "BH-"), ("조사명", ""), ("위치", ""), ("표고(m)", "")]
		self.meta_vars = {}
		for index, (label, default) in enumerate(fields_row1):
			column = index * 2
			ttk.Label(meta, text=label).grid(row=0, column=column, padx=(4, 5), pady=4, sticky="w")
			variable = tk.StringVar(value=default)
			self.meta_vars[label] = variable
			entry = ttk.Entry(meta, textvariable=variable, width=18)
			entry.grid(row=0, column=column + 1, padx=(0, 10), pady=4)
			# 기본 정보의 표고(m)가 변경되면 측점 표고를 자동 업데이트
			if label == "표고(m)":
				variable.trace_add("write", lambda *args: self.recalculate_table_elevations())

		# 2줄: 좌표 X(N), 좌표 Y(E), 최고지하수위 심도, 케이싱심도, SPT 측정종료심도
		fields_row2 = [
			("좌표 X(N)", ""), ("좌표 Y(E)", ""), 
			("최고지하수위 심도(GL-, m)", ""), ("케이싱심도(-m)", ""), ("SPT 측정종료심도(m)", "")
		]
		for index, (label, default) in enumerate(fields_row2):
			column = index * 2
			ttk.Label(meta, text=label).grid(row=1, column=column, padx=(4, 5), pady=4, sticky="w")
			variable = tk.StringVar(value=default)
			self.meta_vars[label] = variable
			ttk.Entry(meta, textvariable=variable, width=14).grid(row=1, column=column + 1, padx=(0, 10), pady=4)

		ttk.Label(meta, text="※ 지하수위 및 케이싱 심도는 GL 기준 깊이(m)로 입력합니다.").grid(row=2, column=0, columnspan=10, sticky="w", padx=4, pady=(2,0))

		toolbar = ttk.Frame(self.input_tab)
		toolbar.pack(fill="x", padx=5, pady=(2, 4))
		ttk.Button(toolbar, text="행 추가", command=self.add_row).pack(side="left", padx=2)
		ttk.Button(toolbar, text="선택 행 삭제", command=self.delete_selected_rows).pack(side="left", padx=2)
		ttk.Button(toolbar, text="CSV 불러오기", command=self.import_csv).pack(side="left", padx=2)
		ttk.Button(toolbar, text="Excel 양식 불러오기", command=self.import_excel).pack(side="left", padx=2)
		ttk.Button(toolbar, text="PDF/이미지 불러오기", command=lambda: self.attach_image(prompt=False)).pack(side="left", padx=2)
		ttk.Button(toolbar, text="SPT/기반암 정리", command=self.organize_spt_rows).pack(side="left", padx=2)
		ttk.Button(toolbar, text="예비평가 실행", command=self.evaluate).pack(side="left", padx=14)
		ttk.Button(toolbar, text="표 복사", command=self.copy_table).pack(side="right", padx=2)
		ttk.Button(toolbar, text="엑셀로 저장", command=self.export_excel).pack(side="right", padx=2)

		quick = ttk.LabelFrame(self.input_tab, text="측점 직접 입력", padding=6)
		quick.pack(fill="x", padx=5, pady=(0, 4))
		self.manual_depth_var = tk.StringVar()
		self.manual_spt_var = tk.StringVar()
		self.manual_soil_var = tk.StringVar()
		ttk.Label(quick, text="심도(m)").pack(side="left", padx=(2, 4))
		ttk.Entry(quick, textvariable=self.manual_depth_var, width=12).pack(side="left", padx=(0, 12))
		ttk.Label(quick, text="SPT-N 표기").pack(side="left", padx=(2, 4))
		tk.Entry(quick, textvariable=self.manual_spt_var, width=14).pack(side="left", padx=(0, 12))
		tk.Label(quick, text="예: 14/30, 50/5").pack(side="left", padx=(0, 12))
		ttk.Label(quick, text="지반 종류(선택)").pack(side="left", padx=(2, 4))
		tk.Entry(quick, textvariable=self.manual_soil_var, width=18).pack(side="left", padx=(0, 12))
		tk.Button(quick, text="측점 추가", command=self.add_manual_row).pack(side="left")

		content = ttk.PanedWindow(self.input_tab, orient="horizontal")
		content.pack(fill="both", expand=True, padx=5, pady=3)
		table_frame = ttk.Frame(content)
		content.add(table_frame, weight=4)
		preview_frame = ttk.LabelFrame(content, text="첨부한 시추주상도", padding=5)
		content.add(preview_frame, weight=2)
		self.manual_preview_canvas = tk.Canvas(preview_frame, background="#f3f5f7", highlightthickness=0)
		self.manual_preview_canvas.pack(fill="both", expand=True)
		self.manual_preview_canvas.create_text(220, 180, text="PDF/이미지를 첨부하면\n여기에 표시됩니다.", fill="#58616b")
		self.manual_preview_canvas.bind("<MouseWheel>", self.zoom_preview)
		self.manual_preview_canvas.bind("<ButtonPress-1>", self.start_preview_drag)
		self.manual_preview_canvas.bind("<B1-Motion>", self.drag_preview)
		
		columns = ("depth", "elevation", "soil", "observation", "spt", "recovery", "sample_no", "method", "form")
		headings = {
			"depth": "심도\n(GL-, m)", "elevation": "표고\n(m)", "soil": "지반 종류", "observation": "현장관찰기록",
			"spt": "SPT-N 표기\n(예: 14/30)", "recovery": "코어회수율\n(%)", "sample_no": "시료번호", "method": "채취방법", "form": "시료형태",
		}
		widths = {"depth": 85, "elevation": 85, "soil": 130, "observation": 300, "spt": 105, "recovery": 90, "sample_no": 85, "method": 90, "form": 100}
		self.tree = ttk.Treeview(table_frame, columns=columns, show="headings", selectmode="extended")
		for column in columns:
			self.tree.heading(column, text=headings[column])
			self.tree.column(column, width=widths[column], minwidth=60, anchor="center")
		y_scroll = ttk.Scrollbar(table_frame, orient="vertical", command=self.tree.yview)
		x_scroll = ttk.Scrollbar(table_frame, orient="horizontal", command=self.tree.xview)
		self.tree.configure(yscrollcommand=y_scroll.set, xscrollcommand=x_scroll.set)
		self.tree.grid(row=0, column=0, sticky="nsew")
		y_scroll.grid(row=0, column=1, sticky="ns")
		x_scroll.grid(row=1, column=0, sticky="ew")
		table_frame.rowconfigure(0, weight=1)
		table_frame.columnconfigure(0, weight=1)
		self.tree.bind("<Double-1>", self.edit_cell)
		self.tree.tag_configure("evenrow", background="#ffffff")
		self.tree.tag_configure("oddrow", background="#f1f5f8")
		ttk.Label(self.input_tab, text="행을 더블클릭하면 수정할 수 있습니다. 심도 변경 시 표고는 지표면 표고 기준으로 자동 수정됩니다.").pack(anchor="w", padx=7, pady=(0, 6))

	def _build_image_tab(self):
		controls = ttk.Frame(self.image_tab, padding=8)
		controls.pack(fill="x")
		ttk.Button(controls, text="시추주상도 첨부 (이미지/PDF)", command=self.attach_image).pack(side="left")
		ttk.Button(controls, text="설치 안내", command=self.show_installation_help).pack(side="left", padx=8)
		self.image_label = ttk.Label(controls, text="첨부된 이미지 없음")
		self.image_label.pack(side="left", padx=12)
		
		self.extract_status = tk.StringVar(value="첨부된 파일 없음")
		ttk.Label(self.image_tab, textvariable=self.extract_status, padding=(12, 4), foreground="#385d7a", font=("맑은 고딕", 10, "bold")).pack(anchor="w")

		paned = ttk.PanedWindow(self.image_tab, orient="horizontal")
		paned.pack(fill="both", expand=True, padx=8, pady=5)
		
		info_frame = ttk.Frame(paned, padding=8)
		paned.add(info_frame, weight=2)
		ttk.Label(info_frame, text="[파일 자동 읽기 안내]\n\n시추주상도 PDF 및 이미지에서 다음 정보를 자동으로 추출합니다:\n- 시추번호, 조사명, 위치, 좌표 X/Y, 표고\n- 케이싱 심도 및 최고지하수위 심도\n- 심도별 SPT-N 표기 및 지반 종류\n\n자동 읽기 완료 후 '시추자료 입력' 탭 및 '예비평가 결과' 탭에서 자동 입력된 결과를 확인하실 수 있습니다.", justify="left", wraplength=400).pack(anchor="nw", pady=10)

		auto_preview_frame = ttk.LabelFrame(paned, text="시추주상도 미리보기", padding=5)
		paned.add(auto_preview_frame, weight=3)
		self.auto_preview_canvas = tk.Canvas(auto_preview_frame, background="#f3f5f7", highlightthickness=0)
		self.auto_preview_canvas.pack(fill="both", expand=True)
		self.auto_preview_canvas.create_text(220, 180, text="PDF/이미지를 불러오면\n여기에 주상도가 표시됩니다.", fill="#58616b")
		self.auto_preview_canvas.bind("<MouseWheel>", self.zoom_preview)
		self.auto_preview_canvas.bind("<ButtonPress-1>", self.start_preview_drag)
		self.auto_preview_canvas.bind("<B1-Motion>", self.drag_preview)

	def _build_result_tab(self):
		self.result_info_var = tk.StringVar(value="조사 정보를 입력하면 여기에 표시됩니다.")
		ttk.Label(self.result_tab, textvariable=self.result_info_var, padding=(10, 8), wraplength=1300).pack(anchor="w")
		self.summary_var = tk.StringVar(value="아직 평가하지 않았습니다.")
		ttk.Label(self.result_tab, textvariable=self.summary_var, style="Result.TLabel", padding=10).pack(anchor="w")
		self.special_notes_var = tk.StringVar(value="현장기록 및 특이사항: 아직 평가하지 않았습니다.")
		ttk.Label(self.result_tab, textvariable=self.special_notes_var, wraplength=1300, justify="left", padding=(10, 0, 10, 8)).pack(anchor="w")
		frame = ttk.Frame(self.result_tab, padding=(10, 0, 10, 10))
		frame.pack(fill="both", expand=True)
		columns = ("depth", "soil", "spt", "n_value", "water", "shallow", "n_limit", "soil_limit", "decision", "reason")
		headings = {
			"depth": "심도(m)", "soil": "지반 종류", "spt": "SPT 표기", "n_value": "N값", "water": "지하수위 하부",
			"shallow": "20m 이내", "n_limit": "N<=25", "soil_limit": "사질토", "decision": "예비평가", "reason": "판정 근거",
		}
		self.result_tree = ttk.Treeview(frame, columns=columns, show="headings")
		for column in columns:
			self.result_tree.heading(column, text=headings[column])
			self.result_tree.column(column, width=110 if column != "reason" else 350, anchor="center")
		y_scroll = ttk.Scrollbar(frame, orient="vertical", command=self.result_tree.yview)
		self.result_tree.configure(yscrollcommand=y_scroll.set)
		self.result_tree.grid(row=0, column=0, sticky="nsew")
		y_scroll.grid(row=0, column=1, sticky="ns")
		frame.rowconfigure(0, weight=1)
		frame.columnconfigure(0, weight=1)
		self.result_tree.tag_configure("target", background="#fff0c2")
		self.result_tree.tag_configure("excluded", background="#eef1f4")
		ttk.Label(self.result_tab, text="예비평가 대상 측점이 존재하더라도, 제외 측점이 하나라도 포함되어 있으면 전체 예비평가는 '예비평가 제외'로 평가됩니다.", wraplength=1250, padding=10).pack(anchor="w")

	def _build_history_tab(self):
		controls = ttk.Frame(self.history_tab, padding=8)
		controls.pack(fill="x")
		ttk.Button(controls, text="선택 기록 보기", command=self.view_history).pack(side="left", padx=2)
		ttk.Button(controls, text="선택한 두 기록 비교", command=self.compare_history).pack(side="left", padx=2)
		ttk.Button(controls, text="비교 결과 저장", command=self.export_history_comparison).pack(side="left", padx=2)
		ttk.Button(controls, text="즐겨찾기 전환", command=self.toggle_history_favorite).pack(side="left", padx=2)
		tk.Button(controls, text="선택 기록 삭제", command=self.delete_history).pack(side="left", padx=2)
		ttk.Button(controls, text="기록 새로고침", command=self.refresh_history_list).pack(side="left", padx=2)
		tk.Label(controls, text="Ctrl 또는 Shift를 누르고 두 기록을 선택하면 비교할 수 있습니다.").pack(side="left", padx=12)
		body = ttk.PanedWindow(self.history_tab, orient="horizontal")
		body.pack(fill="both", expand=True, padx=8, pady=5)
		list_frame = ttk.LabelFrame(body, text="저장된 시추주상도 평가 기록", padding=5)
		body.add(list_frame, weight=1)
		self.history_listbox = tk.Listbox(list_frame, selectmode="extended", exportselection=False, width=38)
		self.history_listbox.pack(side="left", fill="both", expand=True)
		history_scroll = ttk.Scrollbar(list_frame, orient="vertical", command=self.history_listbox.yview)
		history_scroll.pack(side="right", fill="y")
		self.history_listbox.configure(yscrollcommand=history_scroll.set)
		detail_frame = ttk.LabelFrame(body, text="선택 기록의 측점별 평가", padding=5)
		body.add(detail_frame, weight=4)
		columns = ("depth", "soil", "spt", "n_value", "water", "shallow", "n_limit", "soil_limit", "decision", "reason")
		self.history_tree = ttk.Treeview(detail_frame, columns=columns, show="headings")
		for column in columns:
			self.history_tree.heading(column, text=self.result_tree.heading(column, "text"))
			self.history_tree.column(column, width=100 if column != "reason" else 260, anchor="center")
		history_detail_scroll = ttk.Scrollbar(detail_frame, orient="vertical", command=self.history_tree.yview)
		self.history_tree.configure(yscrollcommand=history_detail_scroll.set)
		self.history_tree.pack(side="left", fill="both", expand=True)
		history_detail_scroll.pack(side="right", fill="y")
		self.history_tree.tag_configure("target", background="#fff0c2")
		self.history_tree.tag_configure("excluded", background="#eef1f4")
		self.history_summary = tk.StringVar(value="기록을 선택하세요.")
		ttk.Label(self.history_tab, textvariable=self.history_summary, wraplength=1300, padding=8).pack(anchor="w")
		self.compare_tree = None
		self.refresh_history_list()

	def recalculate_table_elevations(self):
		"""지표면 표고 입력값이 입력되면 측점별 표고(표고 = 지표면표고 - 심도)를 자동 계산합니다."""
		base_elevation = number(self.meta_vars["표고(m)"].get())
		if base_elevation is None:
			return
		for item in self.tree.get_children():
			values = list(self.tree.item(item, "values"))
			depth_val = number(values[0])
			if depth_val is not None:
				values[1] = f"{base_elevation - depth_val:.2f}"
				self.tree.item(item, values=values)

	def load_history(self):
		try:
			with open(self.history_file, "r", encoding="utf-8") as file:
				data = json.load(file)
			return data if isinstance(data, list) else []
		except (OSError, json.JSONDecodeError):
			return []

	def refresh_history_list(self):
		if not hasattr(self, "history_listbox"):
			return
		self.history_listbox.delete(0, "end")
		for index, record in enumerate(self.history):
			favorite = "★ " if record.get("favorite", False) else ""
			partial = "△ " if record.get("evaluation_mode", "일반 예비평가") != "일반 예비평가" else ""
			self.history_listbox.insert("end", f"{favorite}{partial}{index + 1}. {record.get('title', '이전 평가')}")

	def save_history_file(self):
		try:
			with open(self.history_file, "w", encoding="utf-8") as file:
				json.dump(self.history, file, ensure_ascii=False, indent=2)
		except OSError:
			messagebox.showerror("기록 저장 오류", "이전 평가 기록을 저장하지 못했습니다.")

	def toggle_history_favorite(self):
		selected = self.history_listbox.curselection()
		if not selected:
			messagebox.showinfo("기록 선택", "즐겨찾기로 지정할 기록을 선택하세요.")
			return
		for index in selected:
			self.history[index]["favorite"] = not self.history[index].get("favorite", False)
		self.save_history_file()
		self.refresh_history_list()

	def delete_history(self):
		selected = self.history_listbox.curselection()
		if not selected:
			messagebox.showinfo("기록 선택", "삭제할 기록을 선택하세요.")
			return
		if not messagebox.askyesno("기록 삭제", "선택한 이전 평가 기록을 삭제하시겠습니까?", parent=self):
			return
		for index in sorted(selected, reverse=True):
			del self.history[index]
		self.save_history_file()
		self.refresh_history_list()
		for item in self.history_tree.get_children():
			self.history_tree.delete(item)
		self.history_summary.set("기록을 선택하세요.")

	def view_history(self):
		selected = self.history_listbox.curselection()
		if not selected:
			messagebox.showinfo("기록 선택", "먼저 이전 기록을 선택하세요.")
			return
		record = self.history[selected[0]]
		for item in self.history_tree.get_children():
			self.history_tree.delete(item)
		for values in record.get("results", []):
			tag = "target" if len(values) > 8 and values[8] == "예비평가 대상" else "excluded"
			self.history_tree.insert("", "end", values=values, tags=(tag,))
		self.history_summary.set(self.history_summary_text(record))

	def history_summary_text(self, record):
		metadata = record.get("metadata", {})
		targets = record.get("target_depths", [])
		target_text = ", ".join(f"{depth}m" for depth in targets) if targets else "없음"
		return f"{record.get('overall_decision', '예비평가 제외')} | {record.get('evaluation_mode', '일반 예비평가')} | 파일: {record.get('file_name', '수기 입력')} | 조사명: {metadata.get('조사명', '')} | 위치: {metadata.get('위치', '')} | 지하수위: {metadata.get('최고지하수위 심도(GL-, m)', '')}m | 제외 측점: {record.get('excluded_count', 0)}개 | 예비평가 대상 심도: {target_text}"

	def compare_history(self):
		selected = self.history_listbox.curselection()
		if len(selected) != 2:
			messagebox.showinfo("비교 대상 선택", "비교할 기록을 정확히 2개 선택하세요.")
			return
		for child in self.history_tab.winfo_children():
			if getattr(child, "_comparison_frame", False):
				child.destroy()
		compare_frame = ttk.LabelFrame(self.history_tab, text="선택 기록 비교", padding=5)
		compare_frame._comparison_frame = True
		compare_frame.pack(fill="x", padx=8, pady=(0, 8))
		columns = ("record", "file", "survey", "location", "groundwater", "mode", "overall", "targets", "count")
		headings = ("기록", "파일명", "조사명", "위치", "지하수위(m)", "평가 방식", "전체 판정", "예비평가 대상 심도", "대상 개수")
		compare = ttk.Treeview(compare_frame, columns=columns, show="headings", height=3)
		for column, heading in zip(columns, headings):
			compare.heading(column, text=heading)
			compare.column(column, width=150 if column not in ("targets", "survey", "location", "mode", "overall") else 230, anchor="center")
		for number_idx, index in enumerate(selected, start=1):
			record = self.history[index]
			targets = record.get("target_depths", [])
			compare.insert("", "end", values=(number_idx, record.get("file_name", ""), record.get("metadata", {}).get("조사명", ""), record.get("metadata", {}).get("위치", ""), record.get("metadata", {}).get("최고지하수위 심도(GL-, m)", ""), record.get("evaluation_mode", "일반 예비평가"), record.get("overall_decision", "예비평가 제외"), ", ".join(f"{value}m" for value in targets) or "없음", len(targets)))
		compare.pack(fill="x")
		detail_frame = ttk.LabelFrame(self.history_tab, text="지층별·심도별 상세 비교", padding=5)
		detail_frame._comparison_frame = True
		detail_frame.pack(fill="x", padx=8, pady=(0, 8))
		detail_columns = ("depth", "soil_a", "spt_a", "decision_a", "soil_b", "spt_b", "decision_b")
		detail_headings = ("심도(m)", "기록 1 지반", "기록 1 SPT-N", "기록 1 판정", "기록 2 지반", "기록 2 SPT-N", "기록 2 판정")
		detail_tree = ttk.Treeview(detail_frame, columns=detail_columns, show="headings", height=8)
		for column, heading in zip(detail_columns, detail_headings):
			detail_tree.heading(column, text=heading)
			detail_tree.column(column, width=125, anchor="center")
		result_maps = []
		all_depths = set()
		for index in selected:
			mapping = {}
			for values in self.history[index].get("results", []):
				if values:
					depth_value = number(values[0])
					if depth_value is not None:
						mapping[depth_value] = values
						all_depths.add(depth_value)
			result_maps.append(mapping)
		for depth in sorted(all_depths):
			depth_text = f"{depth:g}"
			left = result_maps[0].get(depth, [])
			right = result_maps[1].get(depth, [])
			detail_tree.insert("", "end", values=(depth_text, left[1] if left else "", left[2] if left else "", left[8] if len(left) > 8 else "", right[1] if right else "", right[2] if right else "", right[8] if len(right) > 8 else ""))
		detail_tree.pack(fill="x")

	def add_row(self, values=None):
		default = list(values) if values else ["", "", "", "", "", "", "", "", ""]
		base_elevation = number(self.meta_vars["표고(m)"].get())
		depth_val = number(default[0])
		if base_elevation is not None and depth_val is not None and not str(default[1]).strip():
			default[1] = f"{base_elevation - depth_val:.2f}"
		row_tag = "oddrow" if len(self.rows) % 2 else "evenrow"
		item = self.tree.insert("", "end", values=default, tags=(row_tag,))
		self.rows.append(item)

	def add_manual_row(self):
		depth = self.manual_depth_var.get().strip()
		spt = self.manual_spt_var.get().strip()
		if not depth or not spt:
			messagebox.showwarning("입력 필요", "심도와 SPT-N 표기를 입력하세요.")
			return
		depth_val = number(depth)
		if depth_val is None or parse_spt(spt) is None:
			messagebox.showwarning("입력 형식 확인", "심도는 숫자로, SPT-N은 14/30 또는 50/5 형식으로 입력하세요.")
			return
		base_elevation = number(self.meta_vars["표고(m)"].get())
		elevation_str = f"{base_elevation - depth_val:.2f}" if base_elevation is not None else ""
		self.add_row((depth, elevation_str, self.manual_soil_var.get().strip(), "", spt, "", "", "", ""))
		self.manual_depth_var.set("")
		self.manual_spt_var.set("")
		self.manual_soil_var.set("")

	def delete_selected_rows(self):
		for item in self.tree.selection():
			self.tree.delete(item)
			if item in self.rows:
				self.rows.remove(item)

	def edit_cell(self, event):
		item = self.tree.identify_row(event.y)
		column_id = self.tree.identify_column(event.x)
		if not item or not column_id:
			return
		column_index = int(column_id[1:]) - 1
		bbox = self.tree.bbox(item, column_id)
		if not bbox:
			return
		values = list(self.tree.item(item, "values"))
		editor = ttk.Entry(self.tree)
		editor.insert(0, values[column_index])
		editor.place(x=bbox[0], y=bbox[1], width=bbox[2], height=bbox[3])
		editor.focus_set()

		def finish(_event=None):
			if editor.winfo_exists():
				values[column_index] = editor.get()
				# 심도가 수정되었고 표고가 비어있거나 수정이 일어난 경우 표고 자동 계산
				if column_index == 0:
					depth_val = number(values[0])
					base_elevation = number(self.meta_vars["표고(m)"].get())
					if depth_val is not None and base_elevation is not None:
						values[1] = f"{base_elevation - depth_val:.2f}"
				self.tree.item(item, values=values)
				editor.destroy()

		editor.bind("<Return>", finish)
		editor.bind("<FocusOut>", finish)
		editor.bind("<Escape>", lambda _event: editor.destroy())

	def attach_image(self, prompt=True):
		previous_tab = self.notebook.index(self.notebook.select())
		path = filedialog.askopenfilename(title="시추주상도 이미지 또는 PDF 선택", filetypes=[("이미지/PDF", "*.png *.jpg *.jpeg *.gif *.bmp *.pdf"), ("이미지 파일", "*.png *.jpg *.jpeg *.gif *.bmp"), ("PDF 파일", "*.pdf"), ("모든 파일", "*.*")])
		if not path:
			return
		if self.pdf_document is not None:
			self.pdf_document.close()
			self.pdf_document = None
		self.pdf_page_index = 0
		self.image_path = path
		self.preview_source = None
		self.preview_scale = 1.0
		self.preview_offset = [0, 0]
		for label, variable in self.meta_vars.items():
			variable.set("BH-" if label == "시추번호" and prompt else "")
		self.image_label.configure(text=os.path.basename(path))
		self.extract_status.set("파일 첨부 완료. 자동 읽기를 준비합니다.")
		for item in self.tree.get_children():
			self.tree.delete(item)
		self.rows.clear()
		if path.lower().endswith(".pdf"):
			if fitz is None:
				self.extract_status.set("PDF 텍스트 추출 모듈이 없어 자동 읽기를 진행할 수 없습니다.")
				if prompt:
					self.ask_file_processing(previous_tab)
				else:
					self.show_installation_help(["PyMuPDF"])
				return
			try:
				self.pdf_document = fitz.open(path)
			except (OSError, RuntimeError) as error:
				messagebox.showerror("PDF 열기 오류", str(error))
				return
		
		self.update_idletasks()
		self._redraw_image()
		if prompt:
			self.ask_file_processing(previous_tab)

	def ask_file_processing(self, previous_tab):
		proceed = messagebox.askyesno(
			"시추주상도 정리 및 예비평가",
			"시추주상도 파일의 정리 및 예비평가를 진행하시겠습니까?",
			parent=self,
		)
		if not proceed:
			self.notebook.select(self.input_tab)
			return
		missing = self.missing_dependencies_for_file()
		if missing and self.ask_installation_choice(missing):
			return
		if not self.collect_rows():
			self.extract_document_data()
		if not self.collect_rows():
			messagebox.showinfo("설치 없이 진행", "자동 추출된 측점이 없습니다. 설치하지 않고 진행하려면 CSV 불러오기 또는 입력표 수동 입력을 사용하세요.", parent=self)
			return
		self.organize_spt_rows(show_message=False)
		self.evaluate()

	def ask_installation_choice(self, missing):
		message = (
			"자동 읽기 또는 Excel 기능에 필요한 무료 항목이 없습니다.\n\n"
			f"필요 항목: {', '.join(missing)}\n\n"
			"무료 설치를 진행하시겠습니까?\n"
			"예: 설치 안내창으로 이동\n"
			"아니요: 설치하지 않고 CSV 또는 수동 입력으로 진행"
		)
		install = messagebox.askyesno("무료 설치 여부", message, parent=self)
		if install:
			self.show_installation_help(missing)
		return install

	def missing_dependencies_for_file(self):
		missing = []
		if self.image_path.lower().endswith(".pdf") and fitz is None:
			missing.append("PyMuPDF")
		if not self.image_path.lower().endswith(".pdf"):
			if Image is None:
				missing.append("Pillow")
			tesseract_exists = any(os.path.exists(path) for path in (r"C:\Program Files\Tesseract-OCR\tesseract.exe", r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"))
			if pytesseract is None and not tesseract_exists:
				missing.extend(["pytesseract", "Tesseract OCR"])
		return missing

	def show_installation_help(self, missing=None):
		links = {
			"PyMuPDF": ("PDF 텍스트 읽기", "https://pypi.org/project/PyMuPDF/"),
			"Pillow": ("이미지 파일 읽기", "https://pypi.org/project/Pillow/"),
			"pytesseract": ("Python OCR 연결", "https://pypi.org/project/pytesseract/"),
			"Tesseract OCR": ("이미지 속 숫자·문자 인식 엔진", "https://github.com/UB-Mannheim/tesseract/wiki"),
			"openpyxl": ("Excel 읽기·저장", "https://pypi.org/project/openpyxl/"),
		}
		window = tk.Toplevel(self)
		window.title("설치 안내 및 출력 준비")
		window.geometry("700x470")
		window.transient(self)
		window.grab_set()
		ttk.Label(window, text="자동 읽기·Excel 출력에 필요한 설치 안내", font=("맑은 고딕", 13, "bold"), padding=12).pack(anchor="w")
		message = "필요한 기능은 관련 프로그램 설치 후 사용할 수 있습니다. 아래 링크를 클릭하면 공식 설치 페이지가 열립니다."
		if missing:
			message = "현재 필요한 항목이 설치되지 않았거나 현재 Python 환경에서 인식되지 않았습니다.\n" + ", ".join(missing) + "\n설치 후 프로그램을 다시 실행하세요."
		tk_label = ttk.Label(window, text=message, justify="left", wraplength=650, padding=(12, 0, 12, 10))
		tk_label.pack(anchor="w")
		ttk.Label(window, text="설치 방법", font=("맑은 고딕", 11, "bold"), padding=(12, 2)).pack(anchor="w")
		instructions = (
			"1. 아래에서 'Python 패키지 자동 설치' 버튼을 누릅니다.\n"
			"2. PDF·이미지·Excel에 필요한 무료 패키지가 자동 설치됩니다.\n"
			"3. Tesseract OCR은 아래 공식 링크를 누른 뒤 설치 파일을 내려받습니다.\n"
			"4. 내려받은 설치 파일을 실행하고 Next -> Install -> Finish를 누릅니다.\n"
			"5. 설치가 끝나면 이 프로그램을 닫고 다시 실행합니다. 모든 항목은 무료입니다."
		)
		ttk.Label(window, text=instructions, justify="left", wraplength=650, padding=(12, 2, 12, 8)).pack(anchor="w")
		package_names = {"PyMuPDF": "pymupdf", "Pillow": "pillow", "pytesseract": "pytesseract", "openpyxl": "openpyxl"}
		packages = [package_names[name] for name in (missing or package_names) if name in package_names]
		ttk.Button(window, text="Python 패키지 자동 설치", command=lambda selected=packages: self.install_python_packages_from_app(selected, window)).pack(anchor="w", padx=12, pady=(0, 8))
		ttk.Label(window, text="수동 설치 명령: python -m pip install pymupdf pillow pytesseract openpyxl", wraplength=650, padding=(12, 0, 12, 8)).pack(anchor="w")
		for name, (purpose, url) in links.items():
			row = ttk.Frame(window, padding=(12, 2))
			row.pack(fill="x")
			ttk.Label(row, text=f"{name} - {purpose}", width=34, anchor="w").pack(side="left")
			ttk.Button(row, text="공식 설치 링크 열기", command=lambda address=url: webbrowser.open(address)).pack(side="left")
			ttk.Button(row, text="링크 복사", command=lambda address=url: self.copy_text(address)).pack(side="left", padx=5)
		ttk.Label(window, text="PDF 자동 읽기에는 PyMuPDF, 이미지 자동 읽기에는 Pillow와 Tesseract OCR, Excel 출력에는 openpyxl이 필요합니다.", wraplength=650, padding=12).pack(anchor="w")
		ttk.Button(window, text="닫기", command=window.destroy).pack(pady=10)

	def install_python_packages_from_app(self, packages, window):
		if not packages:
			messagebox.showinfo("설치할 Python 패키지 없음", "현재 Python 패키지는 모두 설치된 것으로 확인됩니다. Tesseract OCR만 별도 설치가 필요할 수 있습니다.", parent=window)
			return
		try:
			process = subprocess.Popen([sys.executable, "-m", "pip", "install", *packages], creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
		except OSError as error:
			messagebox.showerror("자동 설치 오류", str(error), parent=window)
			return
		self.check_installation_process(process, window)

	def check_installation_process(self, process, window):
		if process.poll() is None:
			window.after(500, lambda: self.check_installation_process(process, window))
			return
		if process.returncode == 0:
			messagebox.showinfo("설치 완료", "필요한 Python 패키지 설치가 완료되었습니다. 프로그램을 다시 실행하면 자동 읽기와 Excel 기능을 사용할 수 있습니다.", parent=window)
		else:
			messagebox.showerror("설치 실패", f"설치가 완료되지 않았습니다. 종료 코드: {process.returncode}\n설치 안내창의 수동 명령을 다시 시도하세요.", parent=window)

	def copy_text(self, value):
		self.clipboard_clear()
		self.clipboard_append(value)
		self.update()

	def extract_document_data(self):
		text = ""
		try:
			if self.image_path.lower().endswith(".pdf"):
				if self.pdf_document is None or fitz is None:
					return
				text = "\n".join(page.get_text("text") for page in self.pdf_document)
			else:
				if Image is None:
					return
				image = Image.open(self.image_path)
				if pytesseract is not None:
					try:
						text = pytesseract.image_to_string(image, lang="kor+eng")
					except Exception:
						text = pytesseract.image_to_string(image)
				else:
					executable = next((candidate for candidate in (r"C:\Program Files\Tesseract-OCR\tesseract.exe", r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe") if os.path.exists(candidate)), None)
					if executable is None:
						return
					result = subprocess.run([executable, self.image_path, "stdout", "-l", "eng"], capture_output=True, text=True, encoding="utf-8", errors="ignore", check=False)
					text = result.stdout
		except (OSError, RuntimeError):
			return
		if not text.strip():
			self.extract_status.set("읽을 수 있는 텍스트가 없어 측점 표를 비워 두었습니다.")
			return

		metadata_patterns = {
			"조사명": r"조사명\s+(.+?)(?=\s+발주처|\n|$)",
			"위치": r"위치\s+(.+?)(?=\s+좌표|\n|$)",
			"시추번호": r"시추번호\s+([A-Za-z0-9-]+)",
		}
		for label, pattern in metadata_patterns.items():
			match = re.search(pattern, text, flags=re.IGNORECASE)
			if match and not self.meta_vars[label].get().strip():
				self.meta_vars[label].set(" ".join(match.group(1).split()))

		spt_values = []
		for match in re.findall(r"(?<!\d)\d{1,3}\s*/\s*\d{1,3}(?!\d)", text):
			value = re.sub(r"\s+", "", match)
			if not spt_values or spt_values[-1] != value:
				spt_values.append(value)

		if not spt_values:
			self.extract_status.set("SPT-N 표기가 없어 심도·지층·현장기록만 정리했습니다.")

		# ★ 표고 추출 패턴 강화: 소수점을 포함한 수치 정밀 인식 (예: 22.89)
		elevation_match = re.search(r"표고\s*\(?m\)?\s*[:：]?\s*(\d+(?:\.\d+)?)", text, flags=re.IGNORECASE)
		if not elevation_match:
			elevation_section = text.split("표고", 1)[1].split("시추완료일", 1)[0] if "표고" in text else ""
			elevation_numbers = re.findall(r"\d+(?:\.\d+)?", elevation_section)
			if elevation_numbers:
				elevation_match = re.match(r"(\d+(?:\.\d+)?)", elevation_numbers[0])

		if elevation_match and not self.meta_vars["표고(m)"].get().strip():
			self.meta_vars["표고(m)"].set(elevation_match.group(1))

		water_section = text.split("지하수위", 1)[1].split("사용장비", 1)[0] if "지하수위" in text else ""
		water_numbers = re.findall(r"\d+(?:\.\d+)?", water_section)
		groundwater_match = re.match(r"(\d+(?:\.\d+)?)", water_numbers[-1]) if water_numbers else None
		if groundwater_match and not self.meta_vars["최고지하수위 심도(GL-, m)"].get().strip():
			self.meta_vars["최고지하수위 심도(GL-, m)"].set(groundwater_match.group(1))

		coordinate_section = text.split("좌표", 1)[1].split("표고", 1)[0] if "좌표" in text and "표고" in text else ""
		x_match = re.search(r"X\s*\(?N\)?\s*[:：]?\s*(\d+(?:\.\d+)?)", coordinate_section, flags=re.IGNORECASE)
		y_match = re.search(r"Y\s*\(?E\)?\s*[:：]?\s*(\d+(?:\.\d+)?)", coordinate_section, flags=re.IGNORECASE)
		if x_match and not self.meta_vars["좌표 X(N)"].get().strip():
			self.meta_vars["좌표 X(N)"].set(x_match.group(1))
		if y_match and not self.meta_vars["좌표 Y(E)"].get().strip():
			self.meta_vars["좌표 Y(E)"].set(y_match.group(1))

		casing_section = text.split("케이싱심도", 1)[1].split("지하수위", 1)[0] if "케이싱심도" in text and "지하수위" in text else ""
		casing_numbers = re.findall(r"\d+(?:\.\d+)?", casing_section)
		if casing_numbers and not self.meta_vars["케이싱심도(-m)"].get().strip():
			self.meta_vars["케이싱심도(-m)"].set(casing_numbers[0])

		borehole_section = text.split("굴진심도", 1)[1].split("케이싱심도", 1)[0] if "굴진심도" in text else ""
		borehole_numbers = re.findall(r"\d+(?:\.\d+)?", borehole_section)
		borehole_match = re.match(r"(\d+(?:\.\d+)?)", borehole_numbers[0]) if borehole_numbers else None
		borehole_depth = number(borehole_match.group(1)) if borehole_match else None
		ground_elevation = number(self.meta_vars["표고(m)"].get())
		
		refusal_index = next((index for index, value in enumerate(spt_values) if is_refusal_spt(value)), None)
		soil_count = refusal_index if refusal_index is not None and refusal_index > 0 else len(spt_values)
		interval = infer_spt_interval(borehole_depth, soil_count) if borehole_depth is not None else None
		soil_values = "모래(자동추출)" if "모래" in text else ""
		observation = " ".join(dict.fromkeys(line.strip() for line in text.splitlines() if any(word in line for word in ("모래", "자갈", "점토", "실트", "풍화", "암층"))))
		sample_records = re.findall(r"\b(DS|PS|DC|TC|SC)\s*[-:]?\s*(\d+)\b", text, flags=re.IGNORECASE)
		recovery_values = re.findall(r"(?:코어회수율|회수율)\s*[:：]?\s*(\d{1,3})\s*%", text)
		form_value = ""

		for item in self.tree.get_children():
			self.tree.delete(item)
		self.rows.clear()
		if not spt_values:
			boundary_rows = re.findall(r"(?m)^\s*(\d+(?:\.\d+)?)\s+(\d+(?:\.\d+)?)\s+(\d+(?:\.\d+)?)\s*$", text)
			added_boundary = False
			for depth, elevation, _thickness in boundary_rows:
				if borehole_depth is None or float(depth) <= borehole_depth:
					self.add_row((depth, elevation, "", observation, "", "", "", "", ""))
					added_boundary = True
			if not added_boundary:
				self.add_row(("", "", "", observation, "", "", "", "", ""))
			self.extract_status.set("SPT-N 없음: 심도·지층·현장기록 기준 부분자료 평가를 선택할 수 있습니다.")
			return

		for index, spt in enumerate(spt_values):
			soil = "기반암(자동추출)" if refusal_index is not None and index >= refusal_index else soil_values
			depth = ""
			row_elevation = ""
			if interval is not None:
				depth = f"{interval * (index + 1):.1f}".rstrip("0").rstrip(".")
				if ground_elevation is not None:
					row_elevation = f"{ground_elevation - float(depth):.2f}"
			sample_no = ""
			method = ""
			if index < len(sample_records):
				method = sample_records[index][0].upper()
				sample_no = sample_records[index][1]
			recovery = recovery_values[index] if index < len(recovery_values) else ""
			row_form = form_value if sample_no else ""
			self.add_row((depth, row_elevation, soil, observation, spt, recovery, sample_no, method, row_form))
		if interval is not None:
			self.meta_vars["SPT 측정종료심도(m)"].set(f"{interval * soil_count:.1f}".rstrip("0").rstrip("."))
		self.extract_status.set(f"자동 추출 완료: SPT-N {len(spt_values)}개. 입력표와 예비평가 결과를 확인하세요.")

	def open_external_document(self):
		if not self.image_path:
			messagebox.showwarning("파일 없음", "먼저 이미지 또는 PDF를 첨부하세요.")
			return
		try:
			os.startfile(self.image_path)
		except (AttributeError, OSError) as error:
			messagebox.showerror("파일 열기 오류", str(error))

	def _redraw_image(self, _event=None):
		canvases = [
			getattr(self, "manual_preview_canvas", None),
			getattr(self, "auto_preview_canvas", None)
		]
		for canvas in canvases:
			if canvas is None:
				continue
			canvas.delete("all")
			if not self.image_path:
				canvas.create_text(220, 180, text="시추주상도를 첨부하면\n여기에 표시됩니다.", fill="#58616b")
				continue
			if Image is None or ImageTk is None:
				if self.pdf_document is not None and fitz is not None:
					try:
						page = self.pdf_document.load_page(self.pdf_page_index)
						pixmap = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
						self.preview_image = tk.PhotoImage(data=pixmap.tobytes("png"))
						canvas.create_image(max(canvas.winfo_width() // 2, 150), max(canvas.winfo_height() // 2, 150), image=self.preview_image, anchor="center")
						continue
					except (OSError, RuntimeError, tk.TclError):
						pass
				canvas.create_text(220, 180, text="Pillow가 없어 이미지를 표시할 수 없습니다.\n파일 경로: " + self.image_path)
				continue
			try:
				if self.preview_source is None:
					if self.pdf_document is not None:
						page = self.pdf_document.load_page(self.pdf_page_index)
						pixmap = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
						self.preview_source = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
					else:
						self.preview_source = Image.open(self.image_path).convert("RGB")
				self._draw_preview_on_canvas(canvas)
			except (OSError, tk.TclError) as error:
				canvas.create_text(220, 180, text=f"이미지를 열 수 없습니다: {error}")

	def _draw_preview(self):
		self._redraw_image()

	def _draw_preview_on_canvas(self, canvas):
		if canvas is None or self.preview_source is None:
			return
		canvas.delete("all")
		canvas_width = max(canvas.winfo_width() - 20, 300)
		canvas_height = max(canvas.winfo_height() - 20, 250)
		image_width, image_height = self.preview_source.size
		fit_scale = min(canvas_width / image_width, canvas_height / image_height)
		scale = max(0.1, min(fit_scale * self.preview_scale, 8.0))
		resized = self.preview_source.resize((max(1, int(image_width * scale)), max(1, int(image_height * scale))), Image.Resampling.LANCZOS)
		self.preview_image = ImageTk.PhotoImage(resized)
		center_x = canvas_width // 2 + self.preview_offset[0]
		center_y = canvas_height // 2 + self.preview_offset[1]
		canvas.create_image(center_x, center_y, image=self.preview_image, anchor="center")

	def zoom_preview(self, event):
		if self.preview_source is None:
			return
		self.preview_scale = max(0.25, min(self.preview_scale * (1.15 if event.delta > 0 else 1 / 1.15), 8.0))
		self._draw_preview()

	def start_preview_drag(self, event):
		self.preview_drag_start = (event.x, event.y)

	def drag_preview(self, event):
		if self.preview_drag_start is None:
			return
		last_x, last_y = self.preview_drag_start
		self.preview_offset[0] += event.x - last_x
		self.preview_offset[1] += event.y - last_y
		self.preview_drag_start = (event.x, event.y)
		self._draw_preview()

	def collect_rows(self):
		return [list(self.tree.item(item, "values")) for item in self.tree.get_children() if any(str(value).strip() for value in self.tree.item(item, "values"))]

	def evaluate(self, allow_missing=False):
		groundwater = number(self.meta_vars["최고지하수위 심도(GL-, m)"].get())
		rows = self.collect_rows()
		missing_items = []
		if groundwater is None or groundwater < 0:
			missing_items.append("지하수위")
		if any(number(row[0]) is None for row in rows):
			missing_items.append("심도")
		if any(parse_spt(row[4]) is None for row in rows):
			missing_items.append("SPT-N")
		if missing_items and not allow_missing:
			if "SPT-N" in missing_items and "지하수위" in missing_items:
				fallback_message = "SPT-N 및 지하수위 없이 심도·지반·현장기록 기준으로"
			elif "SPT-N" in missing_items:
				fallback_message = "SPT-N 없이 심도·지반·현장기록 기준으로"
			elif "지하수위" in missing_items:
				fallback_message = "지하수위 없이 심도·SPT-N·지반 기준으로"
			else:
				fallback_message = "누락된 항목을 제외하고 확인 가능한 정보 기준으로"
			message = (
				f"다음 정보가 없거나 불완전합니다: {', '.join(missing_items)}\n\n"
				f"{fallback_message} 예비평가를 진행하시겠습니까?\n"
				"예를 누르면 부분자료 평가로 진행하고, 아니요를 누르면 입력 화면으로 돌아갑니다."
			)
			if messagebox.askyesno("부분자료 예비평가 확인", message, parent=self):
				self.evaluate(allow_missing=True)
			else:
				self.notebook.select(self.input_tab)
			return
		if groundwater is None or groundwater < 0:
			self.evaluation_mode = "지하수위 없이 심도별 진행한 예비평가 결과"
		elif missing_items:
			self.evaluation_mode = "확인 가능한 정보 기준 부분자료 예비평가"
		else:
			self.evaluation_mode = "일반 예비평가"
		for item in self.result_tree.get_children():
			self.result_tree.delete(item)
		file_name = os.path.basename(self.image_path) if self.image_path else "수기 입력"
		water_text = f"{groundwater:g} m" if groundwater is not None and groundwater >= 0 else "자료 없음"
		
		self.result_info_var.set(
			f"{self.evaluation_mode} | 파일: {file_name} | 시추번호: {self.meta_vars['시추번호'].get()} | 조사명: {self.meta_vars['조사명'].get()} | "
			f"위치: {self.meta_vars['위치'].get()} | 좌표: X {self.meta_vars['좌표 X(N)'].get()}, Y {self.meta_vars['좌표 Y(E)'].get()} | "
			f"표고: {self.meta_vars['표고(m)'].get()} m | 케이싱심도: {self.meta_vars['케이싱심도(-m)'].get()} m | 지하수위: {water_text}"
		)
		targets = []
		excluded_count = 0
		special_notes = []
		for row in rows:
			depth = number(row[0])
			n_value = parse_spt(row[4])
			soil_name = str(row[2]).strip().lower()
			observation_text = str(row[3]).strip().lower()
			is_sandy = any(word in soil_name or word in observation_text for word in ("모래", "사질", "sand", "silty"))
			is_rock = any(word in soil_name for word in ("암", "기반암", "rock", "풍화암"))
			below_water = True if groundwater is None or groundwater < 0 else depth is not None and depth >= groundwater
			within_depth = depth is not None and depth <= MAX_EVALUATION_DEPTH
			n_ok = n_value is not None and n_value <= 25
			soil_known = bool(soil_name or observation_text)
			soil_ok = is_sandy and not is_rock
			if not soil_known and self.evaluation_mode != "일반 예비평가":
				soil_ok = True
			target = all((depth is not None, within_depth, n_ok, soil_ok, below_water))
			reasons = []
			if groundwater is None or groundwater < 0:
				reasons.append("지하수위 미확인(심도·SPT-N 기준만 적용)")
			elif not below_water:
				reasons.append("지하수위 상부 지층")
			if not within_depth:
				reasons.append("심도 20m 초과")
			if n_value is None:
				reasons.append("SPT-N 미입력/해석불가")
			elif n_value > 25:
				reasons.append("N값 25 초과")
			if not soil_known:
				reasons.append("지반 종류·현장기록 미확인")
			elif not soil_ok:
				reasons.append("사질토 조건 아님 또는 암반")
			decision = "예비평가 대상" if target else "예비평가 제외"
			water_value = "자료없음" if groundwater is None or groundwater < 0 else "예" if below_water else "아니오"
			item_values = (row[0], row[2], row[4], "-" if n_value is None else f"{n_value:g}", water_value, "예" if within_depth else "아니오", "예" if n_ok else "아니오", "예" if soil_ok else "아니오", decision, "모든 조건 충족" if target and self.evaluation_mode == "일반 예비평가" else ", ".join(reasons))
			self.result_tree.insert("", "end", values=item_values, tags=("target" if target else "excluded",))
			if target:
				targets.append(row[0])
			else:
				excluded_count += 1
			if row[3].strip():
				special_notes.append(f"심도 {row[0]}m: {row[3].strip()}")
		
		overall_decision = "예비평가 제외" if excluded_count > 0 else "예비평가 대상"
		
		if targets:
			interval = ", ".join(f"{depth} m" for depth in targets)
			self.summary_var.set(f"전체 판정: {overall_decision} | {self.evaluation_mode} | 대상 측점: {len(targets)}개 | 제외 측점: {excluded_count}개 | 대상 심도: {interval}")
		else:
			self.summary_var.set(f"전체 판정: {overall_decision} | {self.evaluation_mode} | 대상 측점: 0개 | 제외 측점: {excluded_count}개")
		note_text = "\n".join(special_notes[:20]) if special_notes else "확인된 현장기록 및 특이사항 없음"
		self.special_notes_var.set(f"현장기록 및 특이사항:\n{note_text}")
		self.save_history(targets, overall_decision, excluded_count)
		self.notebook.select(self.result_tab)

	def save_history(self, targets, overall_decision="예비평가 제외", excluded_count=0):
		metadata = {label: variable.get() for label, variable in self.meta_vars.items()}
		results = [list(self.result_tree.item(item, "values")) for item in self.result_tree.get_children()]
		file_name = os.path.basename(self.image_path) if self.image_path else "수기 입력"
		title = f"{metadata.get('시추번호', '')} | {metadata.get('조사명', '')} | {file_name}".strip(" |")
		record = {
			"title": title,
			"file_name": file_name,
			"file_path": self.image_path,
			"metadata": metadata,
			"target_depths": targets,
			"results": results,
			"evaluation_mode": self.evaluation_mode,
			"overall_decision": overall_decision,
			"excluded_count": excluded_count,
		}
		self.history.append(record)
		try:
			with open(self.history_file, "w", encoding="utf-8") as file:
				json.dump(self.history, file, ensure_ascii=False, indent=2)
		except OSError:
			return
		self.refresh_history_list()

	def export_history_comparison(self):
		selected = self.history_listbox.curselection()
		if len(selected) != 2:
			messagebox.showinfo("비교 기록 선택", "'이전 기록/비교' 탭에서 비교할 기록을 정확히 2개 선택하세요.")
			self.notebook.select(self.history_tab)
			return
		if Workbook is None:
			if self.ask_installation_choice(["openpyxl"]):
				return
			messagebox.showinfo("설치 없이 저장", "Excel 저장에는 openpyxl이 필요합니다. 현재 평가표는 '표 복사'로 저장할 수 있습니다.")
			return
		path = filedialog.asksaveasfilename(title="두 평가 기록 비교 저장", defaultextension=".xlsx", filetypes=[("Excel 파일", "*.xlsx")])
		if not path:
			return
		workbook = Workbook()
		comparison_sheet = workbook.active
		comparison_sheet.title = "비교 요약"
		comparison_sheet.append(["항목", "기록 1", "기록 2"])
		records = [self.history[index] for index in selected]
		for label, key in (("파일명", "file_name"), ("조사명", "조사명"), ("위치", "위치"), ("지하수위(m)", "최고지하수위 심도(GL-, m)")):
			values = [record.get("file_name", "") if key == "file_name" else record.get("metadata", {}).get(key, "") for record in records]
			comparison_sheet.append([label, *values])
		comparison_sheet.append(["평가 방식", *[record.get("evaluation_mode", "일반 예비평가") for record in records]])
		comparison_sheet.append(["전체 판정", *[record.get("overall_decision", "예비평가 제외") for record in records]])
		comparison_sheet.append(["제외 측점 수", *[record.get("excluded_count", 0) for record in records]])
		comparison_sheet.append(["예비평가 대상 심도", ", ".join(f"{value}m" for value in records[0].get("target_depths", [])) or "없음", ", ".join(f"{value}m" for value in records[1].get("target_depths", [])) or "없음"])
		comparison_sheet.append(["대상 측점 수", len(records[0].get("target_depths", [])), len(records[1].get("target_depths", []))])
		for number_idx, record in enumerate(records, start=1):
			sheet = workbook.create_sheet(f"기록 {number_idx} 결과")
			sheet.append(["심도", "지반 종류", "SPT 표기", "N값", "지하수위 하부", "20m 이내", "N<=25", "사질토", "예비평가", "판정 근거"])
			for values in record.get("results", []):
				sheet.append(values)
		for sheet in workbook.worksheets:
			for column in sheet.columns:
				sheet.column_dimensions[column[0].column_letter].width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 45)
		try:
			workbook.save(path)
		except OSError as error:
			messagebox.showerror("저장 오류", str(error))
			return
		messagebox.showinfo("저장 완료", f"비교 결과를 저장했습니다.\n{path}")

	def copy_table(self):
		output = io.StringIO()
		writer = csv.writer(output, delimiter="\t", lineterminator="\n")
		writer.writerow([self.tree.heading(column, "text").replace("\n", " ") for column in self.tree["columns"]])
		for item in self.tree.get_children():
			writer.writerow(self.tree.item(item, "values"))
		self.clipboard_clear()
		self.clipboard_append(output.getvalue())
		self.update()
		messagebox.showinfo("복사 완료", "입력표가 탭으로 구분되어 클립보드에 복사되었습니다. 엑셀에 바로 붙여넣을 수 있습니다.")

	def export_excel(self):
		path = filedialog.asksaveasfilename(title="엑셀 파일 저장", defaultextension=".xlsx", filetypes=[("Excel 파일", "*.xlsx")])
		if not path:
			return
		if Workbook is None:
			if self.ask_installation_choice(["openpyxl"]):
				return
			messagebox.showinfo("설치 없이 진행", "Excel 저장 대신 '표 복사' 버튼으로 표를 복사해 사용하세요.")
			return
		workbook = Workbook()
		input_sheet = workbook.active
		input_sheet.title = "시추자료"
		input_sheet.append(["항목", "값"])
		for label, variable in self.meta_vars.items():
			input_sheet.append([label, variable.get()])
		input_sheet.append([])
		input_sheet.append([self.tree.heading(column, "text").replace("\n", " ") for column in self.tree["columns"]])
		for item in self.tree.get_children():
			input_sheet.append(list(self.tree.item(item, "values")))
		result_sheet = workbook.create_sheet("예비평가 결과")
		result_sheet.append([self.result_tree.heading(column, "text").replace("\n", " ") for column in self.result_tree["columns"]])
		for item in self.result_tree.get_children():
			result_sheet.append(list(self.result_tree.item(item, "values")))
		for sheet in workbook.worksheets:
			sheet.freeze_panes = "A2"
			for column in sheet.columns:
				sheet.column_dimensions[column[0].column_letter].width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 10), 45)
		try:
			workbook.save(path)
		except OSError as error:
			messagebox.showerror("저장 오류", str(error))
			return
		messagebox.showinfo("저장 완료", f"엑셀 파일을 저장했습니다.\n{path}")

	def import_excel(self):
		previous_tab = self.notebook.index(self.notebook.select())
		if load_workbook is None:
			if self.ask_installation_choice(["openpyxl"]):
				return
			messagebox.showinfo("설치 없이 진행", "Excel 불러오기 대신 CSV 불러오기 또는 입력표 수동 입력을 사용하세요.")
			return
		path = filedialog.askopenfilename(title="시추자료 Excel 선택", filetypes=[("Excel 파일", "*.xlsx *.xlsm"), ("모든 파일", "*.*")])
		if not path:
			return
		try:
			workbook = load_workbook(path, read_only=True, data_only=True)
		except (OSError, ValueError) as error:
			messagebox.showerror("Excel 불러오기 오류", str(error))
			return
		found = None
		for sheet in workbook.worksheets:
			for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
				texts = [str(value or "").strip().lower() for value in row]
				depth_column = next((index for index, text in enumerate(texts) if "심도" in text), None)
				spt_column = next((index for index, text in enumerate(texts) if "spt" in text or "blows" in text or "n값" in text), None)
				if depth_column is not None and spt_column is not None:
					found = (sheet, row_number, row)
					break
			if found:
				break
		if not found:
			workbook.close()
			messagebox.showerror("열을 찾지 못함", "심도와 SPT-N 열이 함께 있는 표를 찾지 못했습니다.")
			return
		sheet, header_row, header = found
		header_texts = [str(value or "").strip().lower() for value in header]
		column_names = {
			"depth": ("심도",), "elevation": ("표고",), "soil": ("지반명", "지반 종류", "층후"),
			"observation": ("현장관찰", "관찰기록"), "spt": ("spt", "blows", "n값"),
			"recovery": ("회수율",), "sample_no": ("시료번호",), "method": ("채취방법",), "form": ("시료형태",),
		}
		indices = {}
		for name, candidates in column_names.items():
			indices[name] = next((index for index, text in enumerate(header_texts) if any(candidate in text for candidate in candidates)), None)
		new_rows = []
		for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
			values = ["" if indices[name] is None or indices[name] >= len(row) or row[indices[name]] is None else str(row[indices[name]]) for name in column_names]
			if not any(value.strip() for value in values):
				continue
			if not values[0].strip() and not values[4].strip():
				continue
			new_rows.append(values)
		workbook.close()
		for item in self.tree.get_children():
			self.tree.delete(item)
		self.rows.clear()
		for row in new_rows:
			self.add_row(row)
		self.ask_file_processing(previous_tab)

	def organize_spt_rows(self, show_message=True):
		rows = self.collect_rows()
		if not rows:
			messagebox.showwarning("자료 없음", "먼저 Excel 또는 CSV 자료를 불러오거나 입력하세요.")
			return
		end_depth = number(self.meta_vars["SPT 측정종료심도(m)"].get())
		spt_rows = [row for row in rows if str(row[4]).strip()]
		if end_depth is None or end_depth <= 0:
			end_depth = max((number(row[0], 0) or 0 for row in rows), default=0)
		if end_depth <= 0 or not spt_rows:
			messagebox.showerror("정리 기준 부족", "SPT 측정종료심도 또는 SPT-N 값이 필요합니다.")
			return
		refusal_index = next((index for index, row in enumerate(spt_rows) if is_refusal_spt(row[4])), None)
		evaluation_count = refusal_index if refusal_index is not None and refusal_index > 0 else len(spt_rows)
		step = infer_spt_interval(end_depth, evaluation_count)
		base_elevation = number(self.meta_vars["표고(m)"].get())
		spt_index = 0
		for row_item in self.tree.get_children():
			values = list(self.tree.item(row_item, "values"))
			if not str(values[4]).strip():
				continue
			if not str(values[0]).strip():
				depth_val = step * (spt_index + 1)
				values[0] = f"{depth_val:.1f}".rstrip("0").rstrip(".")
				if base_elevation is not None:
					values[1] = f"{base_elevation - depth_val:.2f}"
			if refusal_index is not None and spt_index >= refusal_index:
				values[2] = "기반암"
			self.tree.item(row_item, values=values)
			spt_index += 1
		refusal_text = "관입불가 이후 행을 기반암으로 표시했습니다." if refusal_index is not None else "관입불가 표기가 없어 입력된 지반 종류를 유지했습니다."
		if show_message:
			messagebox.showinfo("SPT 정리 완료", f"측점 {len(spt_rows)}개, 추정 측점 간격 {step:.2f}m\n{refusal_text}")

	def import_csv(self):
		path = filedialog.askopenfilename(title="입력 CSV 선택", filetypes=[("CSV 파일", "*.csv"), ("모든 파일", "*.*")])
		if not path:
			return
		try:
			with open(path, "r", encoding="utf-8-sig", newline="") as file:
				rows = list(csv.reader(file))
		except (OSError, UnicodeError) as error:
			messagebox.showerror("불러오기 오류", str(error))
			return
		if rows and rows[0] and "심도" in rows[0][0]:
			rows = rows[1:]
		for item in self.tree.get_children():
			self.tree.delete(item)
		self.rows.clear()
		for row in rows:
			self.add_row((row + [""] * 9)[:9])
		messagebox.showinfo("불러오기 완료", f"{len(rows)}개 행을 불러왔습니다.")


if __name__ == "__main__":
	app = LiquefactionApp()
	app.mainloop()
